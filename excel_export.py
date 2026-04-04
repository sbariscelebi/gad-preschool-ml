"""
excel_export.py
===============
Saves two Excel workbooks to OUTPUT_DIR.

save_metrics_excel       — CV performance metrics (Table 3 equivalent)
save_actual_pred_excel   — Holdout actual-vs-predicted (Table 4 equivalent)
"""

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils  import get_column_letter

from config import CV_FOLDS, MODEL_SHORT, OUTPUT_DIR

# ── Shared cell styles ─────────────────────────────────────────────────────
_HF   = Font(bold=True, size=12, color="FFFFFF")
_HFI  = PatternFill("solid", fgColor="1D3A5F")
_HA   = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BEST = PatternFill("solid", fgColor="D4EDDA")
_EVEN = PatternFill("solid", fgColor="F1EFE8")
_WHITE = PatternFill("solid", fgColor="FFFFFF")
_ERR  = PatternFill("solid", fgColor="FDECEA")


def _hdr(ws, row, col, val) -> None:
    c = ws.cell(row=row, column=col, value=val)
    c.font      = _HF
    c.fill      = _HFI
    c.alignment = _HA


def _dat(ws, row, col, val, even=False, best=False, fmt=None) -> None:
    c = ws.cell(row=row, column=col, value=val)
    c.font      = Font(bold=best, size=11)
    c.fill      = _BEST if best else (_EVEN if even else _WHITE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    if fmt:
        c.number_format = fmt


# ── Metrics workbook ───────────────────────────────────────────────────────

def save_metrics_excel(results: dict,
                       path: str = None) -> None:
    """Save cross-validation metrics. Best AUC-ROC row highlighted green."""
    if path is None:
        path = os.path.join(OUTPUT_DIR, "GAD_model_metrics.xlsx")

    metrics = ["AUC-ROC", "AP", "F1", "Precision", "Recall", "Accuracy", "MCC"]
    header  = ["Model"] + metrics

    wb = Workbook()
    ws = wb.create_sheet("GAD")

    ws.merge_cells(f"A1:{get_column_letter(len(header))}1")
    c = ws["A1"]
    c.value     = (f"GAD Prediction  —  Model Comparison  "
                   f"({CV_FOLDS}-Fold Stratified CV)")
    c.font      = Font(bold=True, size=14, color="FFFFFF")
    c.fill      = PatternFill("solid", fgColor="0A2540")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for j, h in enumerate(header, 1):
        _hdr(ws, 2, j, h)
    ws.row_dimensions[2].height = 22

    best_model = max(results, key=lambda k: results[k]["AUC-ROC"])

    for ri, (name, mdata) in enumerate(results.items(), 3):
        is_best = (name == best_model)
        c       = ws.cell(row=ri, column=1, value=name)
        c.font  = Font(bold=True, size=11)
        c.fill  = _BEST if is_best else (_EVEN if ri % 2 == 0 else _WHITE)
        c.alignment = Alignment(horizontal="left", vertical="center")

        for j, metric in enumerate(metrics, 2):
            _dat(ws, ri, j, mdata[metric],
                 even=(ri % 2 == 0), best=is_best, fmt="0.0000")
        ws.row_dimensions[ri].height = 20

    ws.column_dimensions["A"].width = 26
    for j in range(2, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.freeze_panes = "A3"

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    wb.save(path)
    print(f"  Saved : {path}")


# ── Actual-vs-predicted workbook ───────────────────────────────────────────

def save_actual_pred_excel(df: pd.DataFrame,
                            path: str = None) -> None:
    """Summary sheet + one detail sheet per model. Errors in light red."""
    if path is None:
        path = os.path.join(OUTPUT_DIR, "GAD_actual_vs_predicted.xlsx")

    wb = Workbook()

    # Summary
    ws0 = wb.create_sheet("Summary", 0)
    ws0.merge_cells("A1:F1")
    c = ws0["A1"]
    c.value     = ("GAD Prediction  —  Actual vs Predicted  "
                   "(Holdout Validation 20%)")
    c.font      = Font(bold=True, size=14, color="FFFFFF")
    c.fill      = PatternFill("solid", fgColor="0A2540")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws0.row_dimensions[1].height = 26

    for j, h in enumerate(
        ["Model", "Total Samples", "Correct", "Incorrect", "Accuracy (%)", "Remark"],
        1,
    ):
        _hdr(ws0, 2, j, h)
    ws0.row_dimensions[2].height = 20

    for ri, (name, grp) in enumerate(df.groupby("Model"), 3):
        total   = len(grp)
        correct = int(grp["Correct"].sum())
        acc     = round(correct / total * 100, 2)
        for j, v in enumerate(
            [name, total, correct, total - correct, acc, "See GAD metrics sheet"],
            1,
        ):
            _dat(ws0, ri, j, v, even=(ri % 2 == 0))
        ws0.row_dimensions[ri].height = 18

    for j, w in enumerate([26, 16, 12, 12, 14, 24], 1):
        ws0.column_dimensions[get_column_letter(j)].width = w
    ws0.freeze_panes = "A3"

    # Per-model sheets
    cols = ["Sample_Index", "Actual", "Predicted", "Probability_GAD", "Correct"]

    for name, grp in df.groupby("Model"):
        ws = wb.create_sheet(title=f"GAD_{MODEL_SHORT[name]}"[:31])

        ws.merge_cells(f"A1:{get_column_letter(len(cols))}1")
        c = ws["A1"]
        c.value     = f"GAD  —  {name}  (Holdout Validation Set)"
        c.font      = Font(bold=True, size=13, color="FFFFFF")
        c.fill      = PatternFill("solid", fgColor="1D3A5F")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        for j, h in enumerate(cols, 1):
            _hdr(ws, 2, j, h)
        ws.row_dimensions[2].height = 20

        for ri, (_, row) in enumerate(grp.iterrows(), 3):
            for j, col in enumerate(cols, 1):
                c = ws.cell(row=ri, column=j, value=row[col])
                c.font      = Font(size=11)
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.fill      = (_ERR if not row["Correct"]
                               else (_EVEN if ri % 2 == 0 else _WHITE))
                if col == "Probability_GAD":
                    c.number_format = "0.0000"
            ws.row_dimensions[ri].height = 18

        for j in range(1, len(cols) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 20
        ws.freeze_panes = "A3"

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    wb.save(path)
    print(f"  Saved : {path}")
